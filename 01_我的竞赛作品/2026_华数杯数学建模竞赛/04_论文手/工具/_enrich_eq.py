# -*- coding: utf-8 -*-
"""为 Q1 问题分析 docx 注入 OMML 原生公式（美化目标函数）"""
import sys
from docx import Document
from docx.oxml import parse_xml

MNS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
WNS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

def r(t):   return f'<m:r><m:t xml:space="preserve">{t}</m:t></m:r>'
def sub(e, s): return f'<m:sSub><m:e>{e}</m:e><m:sub>{s}</m:sub></m:sSub>'
def sup(e, s): return f'<m:sSup><m:e>{e}</m:e><m:sup>{s}</m:sup></m:sSup>'
def frac(n, d): return f'<m:f><m:num>{n}</m:num><m:den>{d}</m:den></m:f>'
def par(e):   return f'<m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr><m:e>{e}</m:e></m:d>'
def om(n):    return f'<m:oMath>{n}</m:oMath>'

def omath_par(nodes):
    xml = (f'<w:p xmlns:w="{WNS}" xmlns:m="{MNS}">'
           f'<w:pPr><w:jc w:val="center"/></w:pPr>{om(nodes)}</w:p>')
    return parse_xml(xml)

def phi(name, body):
    return om(sub(r('Φ'), r(name)) + r('=') + body)

PHI1 = phi('1', frac(r('A'), sub(r('A'), r('norm'))) + r('+') + sup(par(r('R−1')), r('2')))
PHI2 = phi('2', frac(r('max(W,H)'), sub(r('W'), r('norm'))))
CON  = om(r('∀i≠j: ') + par(r('xᵢ≥xⱼ+wⱼ')) + r('∨') + par(r('xⱼ≥xᵢ+wᵢ')) +
         r('∨') + par(r('yᵢ≥yⱼ+hⱼ')) + r('∨') + par(r('yⱼ≥yᵢ+hᵢ')))

def main(docx_path):
    doc = Document(docx_path)
    anchor_intro = anchor_disp = None
    for i, p in enumerate(doc.paragraphs):
        if '针对词典序双目标' in p.text:
            anchor_intro = i
        if '本文采用' in p.text and '四层框架' in p.text:
            anchor_disp = i
    def insert_after(idx, elems):
        a = doc.paragraphs[idx]._p
        for e in elems:
            a.addnext(e)
            a = e
    if anchor_disp is not None:
        insert_after(anchor_disp, [omath_par(CON)])
    if anchor_intro is not None:
        insert_after(anchor_intro, [omath_par(PHI1), omath_par(PHI2)])
    doc.save(docx_path)
    print('enriched OK; anchors intro=%s disp=%s' % (anchor_intro, anchor_disp))

if __name__ == '__main__':
    main(sys.argv[1])
